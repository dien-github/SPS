#!/usr/bin/env python3
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "docs" / "FIRMWARE_SYSTEM_REPORT.md"
OUTPUT = ROOT / "docs" / "Bao_cao_he_thong_firmware_SPS.docx"

ACCENT = RGBColor(31, 78, 121)
BLUE = RGBColor(46, 116, 181)
INK = RGBColor(33, 37, 41)
MUTED = RGBColor(92, 99, 112)
GREEN = RGBColor(0, 112, 60)
RED = RGBColor(180, 35, 24)
PURPLE = RGBColor(112, 48, 160)
CODE_BLUE = RGBColor(0, 92, 171)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "B7C0CC", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_inches: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths_inches[index] * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side, value in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
                item = tc_mar.find(qn(f"w:{side}"))
                if item is None:
                    item = OxmlElement(f"w:{side}")
                    tc_mar.append(item)
                item.set(qn("w:w"), value)
                item.set(qn("w:type"), "dxa")


def add_paragraph_shading(paragraph, fill: str, border: bool = False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        left = borders.find(qn("w:left"))
        if left is None:
            left = OxmlElement("w:left")
            borders.append(left)
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "16")
        left.set(qn("w:space"), "4")
        left.set(qn("w:color"), "2E74B5")


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)
            run.font.color.rgb = PURPLE
        else:
            paragraph.add_run(part)


def code_color(line: str) -> RGBColor:
    stripped = line.strip()
    if not stripped:
        return MUTED
    if "error=0x000C" in line or "error=0x000D" in line or "FAIL" in line:
        return RED
    if "status=0x00" in line or "PASS" in line or "Download verified successfully" in line:
        return GREEN
    if stripped.startswith(("python ", "& ", "$", "cd ", "FT232", "GPIO", "Address")):
        return CODE_BLUE
    if "0x" in line or stripped.startswith(("seq=", "status=", "crc32=", "VTOR")):
        return PURPLE
    return INK


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        paragraph = doc.add_paragraph(style="CodeBlock")
        add_paragraph_shading(paragraph, "F5F7FA", border=True)
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)
        run.font.color.rgb = code_color(line)


def column_widths(rows: list[list[str]]) -> list[float]:
    count = max(len(row) for row in rows)
    total = 6.5
    if count == 1:
        return [total]
    if count == 2:
        return [2.15, 4.35]
    if count == 3:
        return [1.65, 2.45, 2.40]
    if count == 4:
        return [1.20, 1.55, 2.20, 1.55]
    if count == 5:
        return [1.10, 1.20, 1.35, 1.45, 1.40]
    if count == 6:
        return [1.00, 1.00, 1.10, 1.20, 1.20, 1.00]
    return [total / count] * count


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    normalized = [row + [""] * (max_cols - len(row)) for row in rows]
    widths = column_widths(normalized)
    table = doc.add_table(rows=len(normalized), cols=max_cols)
    table.style = "FirmwareGrid"
    set_table_geometry(table, widths)

    for r_idx, row in enumerate(normalized):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            set_cell_borders(cell)
            if r_idx == 0:
                set_cell_shading(cell, "E8EEF5")
            else:
                set_cell_shading(cell, "FFFFFF" if r_idx % 2 else "FAFBFC")

            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if max_cols > 3 and c_idx in (0, max_cols - 1) else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_runs(paragraph, text)
            for run in paragraph.runs:
                run.font.name = "Calibri"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
                run.font.size = Pt(8.5 if max_cols >= 4 else 9.5)
                run.font.color.rgb = INK
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = ACCENT

    doc.add_paragraph()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip()
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells):
            index += 1
            continue
        rows.append(cells)
        index += 1
    return rows, index


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, ACCENT, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = styles.add_style("CodeBlock", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(9)
    code.paragraph_format.left_indent = Inches(0.14)
    code.paragraph_format.right_indent = Inches(0.05)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.05

    table_style = styles.add_style("FirmwareGrid", 3)
    table_style.font.name = "Calibri"
    table_style.font.size = Pt(9)

    header = section.header.paragraphs[0]
    header.text = "Smart Podium System - MCU Firmware Report"
    header.style = styles["Normal"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.text = "SPS PCD | STM32F401 | FreeRTOS | OTA A/B"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(80)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("BÁO CÁO HỆ THỐNG FIRMWARE MCU")
    run.bold = True
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(22)
    run.font.color.rgb = ACCENT

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run("Smart Podium System - Podium Control Device")
    run.font.size = Pt(14)
    run.font.color.rgb = BLUE

    meta = [
        ["Hạng mục", "Thông tin"],
        ["MCU", "STM32F401 Blackpill, detected STM32F401xD/E"],
        ["RTOS", "FreeRTOS"],
        ["Driver policy", "CMSIS / register-level cho code mới"],
        ["Firmware version", "0.2.0"],
        ["Ngày báo cáo", "2026-06-01"],
        ["Trạng thái", "Đã build, nạp, test UART/relay/OTA; còn projector/IR vật lý cần thiết bị bổ sung"],
    ]
    add_markdown_table(doc, meta)

    callout = doc.add_paragraph()
    add_paragraph_shading(callout, "F4F8FC", border=True)
    callout.paragraph_format.left_indent = Inches(0.14)
    callout.paragraph_format.right_indent = Inches(0.10)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(12)
    run = callout.add_run("Tóm tắt: ")
    run.bold = True
    run.font.color.rgb = ACCENT
    callout.add_run(
        "Firmware đã hoạt động trên board thực: UART1 PING/PONG, relay mô phỏng bằng LED, OTA hợp lệ sang slot B, "
        "reject OTA block lỗi và checksum sai. Các phần chưa test vật lý là projector UART2/RS232, IR carrier 38 kHz và relay/motor thật."
    )

    doc.add_page_break()


def build() -> None:
    text = INPUT.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    in_code = False
    code_lines: list[str] = []
    index = 0
    skip_first_title = True

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            rows, next_index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            index = next_index
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text_heading = heading.group(2)
            if level == 1 and skip_first_title:
                skip_first_title = False
                index += 1
                continue
            style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
            doc.add_paragraph(text_heading, style=style)
            index += 1
            continue

        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, stripped[2:])
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            add_inline_runs(paragraph, numbered.group(1))
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline_runs(paragraph, stripped)
        index += 1

    if code_lines:
        add_code_block(doc, code_lines)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
